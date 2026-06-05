import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from openpyxl.utils import get_column_letter

wb = openpyxl.Workbook()

# ═══════════════════════════════════════════
# Sheet 1: 网页设计风格与动态效果汇总
# ═══════════════════════════════════════════
ws = wb.active
ws.title = "设计风格与动态效果"

header_font = Font(name='Microsoft YaHei', size=11, bold=True, color='FFFFFF')
header_fill = PatternFill(start_color='2B1810', end_color='2B1810', fill_type='solid')
cat_font = Font(name='Microsoft YaHei', size=11, bold=True, color='C89B5E')
cat_fill = PatternFill(start_color='F5ECD7', end_color='F5ECD7', fill_type='solid')
normal_font = Font(name='Microsoft YaHei', size=10, color='333333')
link_font = Font(name='Consolas', size=10, color='1a73e8', underline='single')
alt_fill = PatternFill(start_color='FBF7F0', end_color='FBF7F0', fill_type='solid')
thin_border = Border(
    left=Side(style='thin', color='D4AF7A'),
    right=Side(style='thin', color='D4AF7A'),
    top=Side(style='thin', color='D4AF7A'),
    bottom=Side(style='thin', color='D4AF7A'),
)
wrap = Alignment(wrap_text=True, vertical='top')
center = Alignment(horizontal='center', vertical='top', wrap_text=True)

# Title
ws.merge_cells('A1:F1')
ws['A1'] = 'SiteInspire 网页设计风格与动态效果研究汇总'
ws['A1'].font = Font(name='Microsoft YaHei', size=16, bold=True, color='2B1810')
ws['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[1].height = 40

ws.merge_cells('A2:F2')
ws['A2'] = '数据来源: https://www.siteinspire.com | 整理日期: 2026-06-05 | 涵盖 Minimal/Brutalist/Animation/Typographic/Interactive/3D-WebGL 等主流风格'
ws['A2'].font = Font(name='Microsoft YaHei', size=9, color='888888')
ws['A2'].alignment = Alignment(horizontal='center', vertical='center')
ws.row_dimensions[2].height = 22

# Headers
headers = ['序号', '网站名称', '网址', '设计风格分类', '动态效果 / 交互技术', '技术栈']
widths = [6, 24, 44, 30, 48, 36]
for col, (h, w) in enumerate(zip(headers, widths), 1):
    cell = ws.cell(row=4, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
    cell.border = thin_border
    ws.column_dimensions[get_column_letter(col)].width = w
ws.row_dimensions[4].height = 30

# Full data
sites = [
    ['CATEGORY', '▎极简主义 + 字体排印 (Minimal + Typographic) — siteInspire ~731 极简 / ~2,185 字体排印', '', '', '', ''],
    ['1', 'MONO Design System', 'https://mono.layogtima.com', '极简 / 字体排印 / 单色',
     '纯字体驱动布局，无任何图像；hover 微妙过渡；黑白灰单色调；Space Mono 字体',
     'Tailwind + Space Mono'],
    ['2', 'Francesco Zorzi', 'https://www.francescozorzi.me', '极简 / 终端风格 / 字体排印',
     '终端命令行视觉风格；Tokyo Night 暗色调色板；monospace 字体主导；极简交互',
     '自定义 CSS + Tokyo Night Palette'],
    ['3', 'Charlie Osborn Portfolio', 'https://charliedesignbyans.framer.website', '极简 / 编辑风 / 字体排印',
     '编辑式排版；平衡的留白间距；微妙过渡动画；昼夜模式',
     'Framer'],
    ['4', 'TGW Minimal Studio', 'https://thegoodworks.framer.website', '极简 / 期刊风 / 人文设计',
     '期刊式排版；轻量 hover 过渡；人文主义字体；安静克制的设计语言',
     'Framer'],
    ['5', 'Bruno Tomé', 'https://www.brunotome.com', '极简 / 非对称 / 倒色交互',
     '反转色 hover 动画（黑白切换）；视差滚动；非对称网格布局；单页设计',
     '自定义开发 + GSAP'],

    ['CATEGORY', '▎粗野主义 / 新粗野主义 (Brutalist / Neo-Brutalist) — siteInspire Unusual Layout ~628', '', '', '', ''],
    ['6', 'Eloy Benoffi (Eloyb)', 'https://www.eloyb.design', '粗野主义 / 故障艺术 / 实验性',
     'GSAP ScrollTrigger 驱动故障动画；ScrambleText 文字扰乱特效；Draggable 拖拽交互；ASCII 艺术元素；像素级故障着色器；Awwwards Honorable Mention',
     'Webflow + GSAP (ScrollTrigger/SplitText/ScrambleText/Draggable)'],
    ['7', 'Problem Studio', 'https://problem.studio', '新粗野主义 / 趣味交互 / 非常规导航',
     '整蛊坐垫交互——点击弹射穿越 emoji 进入隐藏联系区；破坏性网格；块状色彩碰撞；出人意料的趣味体验',
     'Webflow + 自定义 JS'],
    ['8', 'OXYZ3', 'https://www.oxyz3.com', '未来粗野 / 反乌托邦 / 中国艺术科技',
     '3D 元素+粗野排版融合；暗黑工业风氛围；超大字号视觉冲击；未来主义反乌托邦美学',
     'Three.js + WebGL'],
    ['9', 'Lydia Amaruch', 'https://www.lydiaamaruch.com', '新粗野主义 / 暗色 / 柏林设计',
     '暗色 UI 新粗野主义；大胆排版；干净克制动效；交互式作品集网格',
     'Webflow / Framer'],
    ['10', 'Stefan Vitasovic', 'https://www.stefanvitasovic.com', '粗野主义 / 故障效果 / 开发者风',
     '失真效果+刻意故障；开发者作品集；高对比度排版；混沌美学',
     '自定义开发 + GSAP'],

    ['CATEGORY', '▎非常规布局 + 动画驱动 (Unusual Layout + Animation) — siteInspire Web & Interactive ~853', '', '', '', ''],
    ['11', 'Roze Bunker', 'https://www.rozebunker.com', '非常规布局 / 艺术日志 / 趣味动画',
     '破坏性网格布局；趣味交互动画；环保主题与粗野风格的冲突融合；艺术日志+电商结合',
     '自定义开发'],
    ['12', 'Design Thinkers 2025', 'https://www.designthinkers2025.com', '会议活动 / 失真效果 / 大胆设计',
     '复杂失真效果；超大字体占据画面；刻意破坏的视觉逻辑；滚动驱动动画',
     'Webflow + GSAP'],
    ['13', 'Client Meeting 2025', 'https://www.clientmeeting2025.jp', '非常规导航 / 横向滚动 / 日本设计',
     '右到左横向滚动（日本特有）；超大动画元素；日式排版美学；企业活动页',
     '自定义开发'],
    ['14', 'Even Odd', 'https://www.evenodd.studio', '创意工作室 / 视频层叠 / 重度动画',
     '多层视频堆叠叠加；大量滚动驱动动画；沉浸式视觉叙事；创意工作室调性',
     '自定义开发 + 视频处理'],
    ['15', '18 Years of Büro', 'https://www.burocratik.com', '粗野时间线 / 机构 / 叙事',
     '粗野主义时间线叙事；庆祝18周年的实验性视觉；大胆布局',
     '自定义开发'],

    ['CATEGORY', '▎交互体验 + 3D/WebGL (Interactive + 3D/WebGL) — 前沿技术驱动', '', '', '', ''],
    ['16', 'Phantom.land', 'https://www.phantom.land', '3D粒子 / WebGL着色器 / 后处理',
     '~78,400粒子生成3D人面轮播；自定义GLSL网格扭曲着色器；拖拽缩放+弹簧物理惯性；卷曲噪声动画；鼠标环境光偏移；景深后处理；GSAP面部过渡',
     'React Three Fiber + GLSL Shaders + GSAP + Next.js'],
    ['17', 'ProofLabs (太空奥德赛)', 'https://www.prooflabs.space', '3D粒子地球 / 太空叙事 / Framer',
     '自定义3D交互粒子地球（粘性滚动）；环绕光标（hover变卫星）；自定义加载动画（地球→轨道→着陆）；滚动触发功能揭示+多层视差',
     'Framer + 3D 粒子系统'],
    ['18', 'ModernStillness (3D壁炉)', 'https://modernstillness.com', '3D环境 / 实时模拟 / 电影级',
     '实时火焰模拟+发光渲染；GPU InstancedMesh落叶物理（风/生命周期/淡出）；尘埃粒子浮动；自适应性能管理（动态pixelRatio/FPS调整）',
     'Three.js + GSAP + Vanilla JS'],
    ['19', '3D Interactive Gallery', 'https://github.com/abx15/3d-interactive-gallery', '3D画廊 / WebGL / 开源',
     'Three.js 网格变形实现图像深度；滚动视差+hover RGB偏移+缩放；自定义顶点/片段着色器；触摸+鼠标拖拽；开源可学习',
     'Three.js + GSAP + WebGL (开源)'],

    ['CATEGORY', '▎多彩配色 + 网格布局 (Colourful + Grid Layout) — siteInspire ~634 Grid Layout', '', '', '', ''],
    ['20', 'DONÜTS by Will Robinson', 'https://www.donuts.framer.website', '波普艺术 / 多彩 / 单页叙事',
     '大胆波普配色；滚动叙事驱动；弹窗式内容揭示；趣味微交互；高饱和色彩碰撞',
     'Framer'],
    ['21', 'Superlist', 'https://www.superlist.com', '产品展示 / 配色切换 / 极简多彩',
     '配色方案滑块动态切换工作vs个人场景；视差滚动；微交互细节打磨到位',
     '自定义开发'],
    ['22', 'Moon UltraLight', 'https://www.moonultralight.com', '产品页 / Apple风格 / 极简多彩',
     'Apple式极简产品页；视差滚动+微交互；干净产品摄影+精致排版',
     '自定义开发'],

    ['CATEGORY', '▎编辑风 + 字体排印主导 (Editorial + Typographic) — siteInspire ~2,185 Typographic', '', '', '', ''],
    ['23', 'Chiara Luzzana', 'https://www.chiaraluzzana.com', '实验性 / 视听融合 / 字体排印',
     '字体作为核心布局元素；音频驱动视觉变化；实验性导航；视听融合体验',
     '自定义开发 + Web Audio API'],
    ['24', 'Pierre-Louis Labonne', 'https://www.pierrelouislabonne.com', '怀旧卡通 / 叙事驱动 / Webflow决赛',
     '怀旧卡通风格插画；滚动叙事驱动；趣味交互动画；Webflow Awards 决赛入围',
     'Webflow'],
    ['25', 'The Goonies', 'https://www.thegoonies.com', '沉浸叙事 / 复古 / 电影级滚动',
     'Awwwards Site of the Day；沉浸式复古滚动叙事；电影级视效体验',
     '自定义开发 + 视频/动画'],

    ['CATEGORY', '▎暗色氛围 / 情绪化设计 (Dark Mode / Moody) — 2025 主流趋势', '', '', '', ''],
    ['26', 'godaylight.com', 'https://godaylight.com', '暗色 / 毛玻璃 / 暖色点缀 / 全屏',
     'Lenis平滑滚动；多层视差；frosted glass毛玻璃卡片+鼠标光晕追踪；大号衬线标题逐行滑入；滚动驱动背景交叉淡入淡出；sticky横向滚动卡片；自定义光标',
     'Next.js + Lenis + 自定义CSS'],
    ['27', 'SUBBIO Studios', 'https://www.subbio.com', '暗色 / 超大字体 / 隐藏交互',
     '超大字体布局；hover显示隐藏画廊；故意破坏的网格系统；暗色沉浸氛围',
     '自定义开发'],
    ['28', 'Curry Cafe', 'https://www.currycafe.com', '旧派混沌 / 现代美学 / 暗色+块色',
     '旧派混沌设计+现代美学融合；强块色碰撞；故意破坏的网格；个性表达',
     '自定义开发'],

    ['CATEGORY', '▎电商 / 品牌 (E-Commerce / Brand) — siteInspire ~797 E-Commerce', '', '', '', ''],
    ['29', 'Bestsized', 'https://www.bestsized.com', '复古菜单 / 品牌叙事 / 质感',
     '复古餐厅菜单灵感；结构混沌但纹理深度到位；品牌故事叙事驱动',
     '自定义开发'],
    ['30', 'Magram', 'https://www.magram.com', '极简新粗野 / 品牌 / 失真',
     '极简版新粗野主义；滚动驱动失真效果；品牌克制表达',
     '自定义开发'],
    ['31', 'Gumroad', 'https://www.gumroad.com', '新粗野主义 / 电商 / 霓虹',
     '完整新粗野主义游乐场；霓虹调色板；对抗性超大字体；锯齿边框元素',
     '自定义开发'],

    ['CATEGORY', '▎创意机构 / 作品集 (Agency / Portfolio) — siteInspire ~1,790 Portfolio', '', '', '', ''],
    ['32', 'Zunc Studio', 'https://www.zuncstudio.com', '极简粗野 / 实验交互 / 印尼创意',
     '极简粗野+实验交互融合；光标特效驱动体验；印尼创意工作室调性',
     '自定义开发'],
    ['33', 'E—B Agency', 'https://www.e-b.agency', '粗野布局 / 品牌机构 / 超大字体',
     '混沌布局；超大字体品牌表达；重度动画驱动页面',
     '自定义开发 + GSAP'],
    ['34', 'Studio Brot', 'https://www.studiobrot.com', '极简 / 光标交互 / 德国创意',
     '极简结构+迷人光标特效；hover内容揭示；德国创意机构严谨风格',
     '自定义开发'],
    ['35', 'Michael Brown', 'https://www.michaelbrown.design', '故障艺术 / 微交互 / 作品集',
     '大量微交互细节；刻意故障效果；混沌hero排版；个人特色强烈',
     '自定义开发 + GSAP'],
    ['36', 'Tomoya Okada (Portfolio v7)', 'https://www.tomoyaokada.com', '实验性作品集 / 日本 / 前端',
     '日本前端工程师实验性作品集；Orpetron六月2025推荐；独特交互设计',
     '自定义开发'],
]

row = 5
for i, site in enumerate(sites):
    is_cat = site[0] == 'CATEGORY'
    if is_cat:
        ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=6)
        cell = ws.cell(row=row, column=1, value=site[1])
        cell.font = cat_font
        cell.fill = cat_fill
        cell.alignment = Alignment(vertical='center', wrap_text=True)
        cell.border = thin_border
        # Apply border to merged cells
        for c in range(2, 7):
            ws.cell(row=row, column=c).border = thin_border
    else:
        for col, val in enumerate(site, 1):
            cell = ws.cell(row=row, column=col, value=str(val))
            cell.border = thin_border
            cell.font = normal_font
            cell.alignment = wrap
            if col == 1:
                cell.alignment = center
            if col == 3:
                cell.font = link_font
            if (i - sum(1 for s in sites[:i] if s[0]=='CATEGORY')) % 2 == 0:
                cell.fill = alt_fill
    ws.row_dimensions[row].height = 65 if site[0] != 'CATEGORY' else 30
    row += 1

# ═══════════════════════════════════════════
# Sheet 2: 设计风格分类体系
# ═══════════════════════════════════════════
ws2 = wb.create_sheet("风格分类体系")
ws2.merge_cells('A1:D1')
ws2['A1'] = 'SiteInspire 设计风格分类体系（Taxonomy）'
ws2['A1'].font = Font(name='Microsoft YaHei', size=14, bold=True, color='2B1810')
ws2['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws2.row_dimensions[1].height = 35

for col, (h, w) in enumerate(zip(['分类维度', '类别名称', '收录量', '设计说明'], [16, 26, 10, 60]), 1):
    cell = ws2.cell(row=3, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = thin_border
    ws2.column_dimensions[get_column_letter(col)].width = w

cats = [
    ['设计风格', 'Minimal 极简主义', '731+', '大量留白、克制配色、减少装饰、内容优先、负空间运用'],
    ['设计风格', 'Typographic 字体排印', '2,185+', '字体作为核心视觉元素，大字号、粗细对比、字体实验、可变字体'],
    ['设计风格', 'Grid Layout 网格布局', '634+', '严格网格系统、模块化卡片、结构化信息层级'],
    ['设计风格', 'Unusual Layout 非常规布局', '628+', '打破传统网格、非对称、重叠元素、破坏边界'],
    ['设计风格', 'Colourful 多彩', '—', '大胆配色、渐变碰撞、高饱和、活力视觉语言'],
    ['设计风格', 'Greyscale 灰度', '—', '黑白灰单色调、强调形状轮廓和排版节奏'],
    ['设计风格', 'Illustrative 插画风', '—', '定制插画为主要视觉语言、品牌个性表达'],
    ['设计风格', 'Horizontal Layout 横向', '—', '横向滚动为主的页面结构、打破纵向阅读习惯'],
    ['网站类型', 'Portfolio 作品集', '1,790+', '个人/团队作品展示、创意行业最常见类型'],
    ['网站类型', 'Design & Art Direction', '2,052+', '强调视觉概念和艺术表达、前卫设计探索'],
    ['网站类型', 'Web & Interactive Design', '853+', '以交互体验为核心的实验性网站'],
    ['网站类型', 'E-Commerce 电商', '797+', '在线商店、产品展示与购买流程体验设计'],
    ['网站类型', 'Editorial 编辑出版', '—', '文章、杂志、新闻内容排版、长文阅读体验'],
    ['网站类型', 'Agencies 创意机构', '—', '创意机构企业展示、品牌形象网站'],
    ['网站类型', 'Fashion 时尚', '731+', '时尚品牌、服装、奢侈品行业视觉设计'],
    ['主题领域', 'Art 艺术', '448+', '画廊、艺术家、博物馆等文化艺术类'],
    ['主题领域', 'Photography 摄影', '—', '摄影师作品集、以图片展示为核心'],
    ['主题领域', 'Culture 文化', '—', '文化机构、活动、教育类'],
    ['主题领域', 'Music 音乐', '—', '音乐人、厂牌、音乐平台'],
    ['动效交互', 'Use of Animation 动画', '458+', '滚动动画、hover微交互、页面过渡、逐帧动画'],
    ['动效交互', 'Unusual Navigation 非常规导航', '—', '非传统菜单、隐藏式导航、实验性导航模式'],
    ['动效交互', 'Copywriting 文案驱动', '—', '以文案为核心叙事手段的设计'],
]

row2 = 4
for i, (dim, cat, cnt, desc) in enumerate(cats):
    for col, val in enumerate([dim, cat, cnt, desc], 1):
        cell = ws2.cell(row=row2, column=col, value=str(val))
        cell.font = normal_font
        cell.border = thin_border
        cell.alignment = wrap
        if i % 2 == 0:
            cell.fill = alt_fill
    ws2.row_dimensions[row2].height = 28
    row2 += 1

# ═══════════════════════════════════════════
# Sheet 3: 2025-2026 设计趋势
# ═══════════════════════════════════════════
ws3 = wb.create_sheet("设计趋势2025-2026")
ws3.merge_cells('A1:C1')
ws3['A1'] = '2025-2026 Web 设计趋势与技术方向'
ws3['A1'].font = Font(name='Microsoft YaHei', size=14, bold=True, color='2B1810')
ws3['A1'].alignment = Alignment(horizontal='center', vertical='center')
ws3.row_dimensions[1].height = 35

for col, (h, w) in enumerate(zip(['设计趋势', '详细描述', '关键技术'], [30, 60, 38]), 1):
    cell = ws3.cell(row=3, column=col, value=h)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = Alignment(horizontal='center', vertical='center')
    cell.border = thin_border
    ws3.column_dimensions[get_column_letter(col)].width = w

trends = [
    ['深度粒子系统', '从2D照片+深度图生成数万粒子构成3D人面/物体，实现轻量级3D展示（纹理仅15KB）。粒子随鼠标/滚动产生有机运动', 'React Three Fiber + GLSL Shaders + GSAP'],
    ['后处理着色器扭曲', '自定义GLSL着色器实现镜头畸变、流体扭曲、晕影、色差等电影级后期特效，大幅提升视觉质感', 'WebGL Shaders + Three.js Post-Processing'],
    ['6D / 多维感官体验', '超越传统3D维度，融合空间音效、物理引擎、粒子系统、触觉反馈打造全感官沉浸式体验', 'Three.js + Web Audio API + Matter.js'],
    ['磁性光标 + 粒子拖尾', '光标与页面元素产生磁力吸附效应，hover时形态变换（箭头→圆环→图标）并生成粒子轨迹', 'GSAP + Canvas 2D + 自定义JS'],
    ['滚动驱动电影级叙事', 'ScrollTrigger驱动的逐帧动画，页面滚动=电影播放进度控制。配合视频序列帧实现丝滑叙事', 'GSAP ScrollTrigger + Lenis + 视频序列帧'],
    ['新粗野主义 (Neo-Brutalism)', '刻意破坏网格、超大字体、块色碰撞、故障效果、锯齿边框、反精致审美运动', 'Webflow/Framer + GSAP + 自定义CSS'],
    ['毛玻璃 + 暖色点缀', 'backdrop-filter毛玻璃卡片+金色/暖色点缀+深色底色，打造高级质感。鼠标追踪光晕增强互动感', 'CSS backdrop-filter + Lenis平滑滚动'],
    ['自适应性能管理', '根据设备GPU/CPU能力动态调整pixelRatio和FPS目标，保证低端设备流畅体验同时高端设备画质拉满', 'Three.js Performance Monitor + 自适应逻辑'],
    ['AI辅助个性化UI', '基于用户行为的实时内容调整、智能配色推荐、生成式UI组件、对话式界面', 'AI API + React/Vue动态组件'],
    ['非常规导航模式', '横向滚动、隐藏菜单、光标驱动导航、3D空间导航、手势交互，打破传统汉堡菜单范式', 'GSAP + CSS Transforms + WebGL'],
    ['字体即设计 (Type-as-Design)', '字体排印作为唯一/核心视觉元素，无图像或少图像的纯字体设计。可变字体+动画字体兴起', 'Variable Fonts + CSS Animations + 自定义字体'],
    ['暗色沉浸 + 生物荧光', '暗色模式基底+霓虹发光点缀+玻璃态深度层次+3D倾斜卡片。暗而不闷、有呼吸感的暗色设计', 'CSS Custom Properties + Three.js + Framer Motion'],
    ['破坏性/实验性布局', '刻意违反设计规则：重叠文本、断裂网格、溢出元素、随机定位。在混乱中寻找秩序', 'CSS Grid/Flexbox 创意用法 + JS动态定位'],
    ['复古未来主义', '结合80-90年代视觉元素(像素、CRT扫描线、霓虹)与未来感技术(3D、粒子)的混搭美学', 'CSS Filters + Three.js + 自定义着色器'],
]

row3 = 4
for i, (trend, desc, tech) in enumerate(trends):
    for col, val in enumerate([trend, desc, tech], 1):
        cell = ws3.cell(row=row3, column=col, value=str(val))
        cell.font = normal_font
        cell.border = thin_border
        cell.alignment = wrap
        if i % 2 == 0:
            cell.fill = alt_fill
    ws3.row_dimensions[row3].height = 55
    row3 += 1

# Save
filepath = r'D:\Cursor\siteinspire-design-analysis.xlsx'
wb.save(filepath)
print(f'✅ Saved: {filepath}')
print(f'   Sheet 1 "{ws.title}": {ws.max_row} 行 (36个网站 + 分类标题)')
print(f'   Sheet 2 "{ws2.title}": {ws2.max_row} 行 (22个分类)')
print(f'   Sheet 3 "{ws3.title}": {ws3.max_row} 行 (14个趋势)')
